"""Acquire, transform and validate the 2026 US elections dashboard datasets.

The whole data pipeline lives in this module so its lineage is inspectable:

* ``acquire()`` downloads the raw sources into ``data/raw/`` (live mode only);
* ``transform_polls()`` turns FiveThirtyEight's pollster-ratings database into one
  record per poll question with Democratic-minus-Republican margins and errors;
* ``transform_seats()`` lists the Senate seats on the 2026 ballot from the
  current-legislator roster and attaches each seat's last result and the state's
  2024 presidential margin;
* ``derive_*()`` recompute every derived column from base columns;
* ``validate()`` checks every dataset against ``dashboard.yml`` and stops on failure.

Offline mode (``DASHBOARD_SAMPLE_MODE=1``) reads the bundled samples, recomputes
the derived columns, validates and writes only the declared processed outputs.

Usage::

    DASHBOARD_SAMPLE_MODE=1 python -m src.data      # offline rebuild (used by the checker)
    python -m src.data --acquire                    # refresh data/raw/ from the live sources
    python -m src.data                              # live transform from data/raw/
    python -m src.data --refresh-sample             # live transform, then rewrite the samples
    python -m src.data --validate                   # validate the processed outputs
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path
from urllib.error import URLError
from urllib.request import urlopen

import pandas as pd
import yaml

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
RAW = DATA / "raw"
MANIFEST = ROOT / "dashboard.yml"

RAW_SOURCES = {
    "raw_polls.csv": "https://raw.githubusercontent.com/fivethirtyeight/data/master/pollster-ratings/raw_polls.csv",
    "election_results_senate.csv":
        "https://raw.githubusercontent.com/fivethirtyeight/election-results/main/election_results_senate.csv",
    "election_results_presidential.csv":
        "https://raw.githubusercontent.com/fivethirtyeight/election-results/main/election_results_presidential.csv",
    "legislators-current.yaml":
        "https://raw.githubusercontent.com/unitedstates/congress-legislators/main/legislators-current.yaml",
}

RACE_TYPES = {"Sen-G": "Senate", "Gov-G": "Governor", "House-G": "House", "Pres-G": "President"}
SEAT_CLASS = {1: "Class I", 2: "Class II", 3: "Class III"}
STATE_NAMES = {
    "AK": "Alaska", "AL": "Alabama", "AR": "Arkansas", "CO": "Colorado", "DE": "Delaware", "FL": "Florida",
    "GA": "Georgia", "IA": "Iowa", "ID": "Idaho", "IL": "Illinois", "KS": "Kansas", "KY": "Kentucky",
    "LA": "Louisiana", "MA": "Massachusetts", "ME": "Maine", "MI": "Michigan", "MN": "Minnesota",
    "MS": "Mississippi", "MT": "Montana", "NC": "North Carolina", "NE": "Nebraska", "NH": "New Hampshire",
    "NJ": "New Jersey", "NM": "New Mexico", "OH": "Ohio", "OK": "Oklahoma", "OR": "Oregon", "RI": "Rhode Island",
    "SC": "South Carolina", "SD": "South Dakota", "TN": "Tennessee", "TX": "Texas", "VA": "Virginia",
    "WV": "West Virginia", "WY": "Wyoming",
}
ELECTION_DAY_2026 = "2026-11-03"

DATASET_FILES = {
    "polls": ("data/sample.csv", "data/processed.csv"),
    "seats": ("data/sample_seats.csv", "data/processed_seats.csv"),
    "swing": ("data/sample_swing.csv", "data/processed_swing.csv"),
    "polls_2026": ("data/sample_polls_2026.csv", "data/processed_polls_2026.csv"),
    "generic_2026": ("data/sample_generic_2026.csv", "data/processed_generic_2026.csv"),
}

# Curated source: a spreadsheet maintained by hand, one row per poll, each with its source URL.
CURATED_XLSX = DATA / "curated" / "2026_midterms.xlsx"


# --------------------------------------------------------------------------- acquire
def acquire() -> None:
    """Download every raw source into data/raw/ with a timeout and clear errors."""
    RAW.mkdir(parents=True, exist_ok=True)
    for name, url in RAW_SOURCES.items():
        try:
            with urlopen(url, timeout=60) as response:
                body = response.read()
        except (URLError, TimeoutError, OSError) as error:
            sys.exit(
                f"Could not fetch {url}: {error}. Check your connection, or run "
                "offline with DASHBOARD_SAMPLE_MODE=1 to rebuild from the bundled samples."
            )
        if not body.strip():
            sys.exit(f"{url} returned an empty body; refusing to overwrite {name}")
        (RAW / name).write_bytes(body)
        print(f"Wrote {RAW / name} ({len(body):,} bytes)")


# --------------------------------------------------------------------------- transform
def transform_polls() -> pd.DataFrame:
    """One row per poll question in a Democrat-vs-Republican general election, 1998-2022."""
    raw = pd.read_csv(RAW / "raw_polls.csv", low_memory=False)
    keep = raw.loc[
        raw["type_simple"].isin(list(RACE_TYPES))
        & (raw["cand1_party"] == "DEM")
        & (raw["cand2_party"] == "REP")
        & (raw["cycle"] % 2 == 0)
    ].copy()
    return pd.DataFrame(
        {
            "question_id": keep["question_id"].astype(int),
            "poll_id": keep["poll_id"].astype(int),
            "cycle": keep["cycle"].astype(int),
            "race_type": keep["type_simple"].map(RACE_TYPES),
            "race": keep["race"],
            "location": keep["location"],
            "pollster": keep["pollster"],
            "methodology": keep["methodology"],
            "partisan_sponsor": keep["partisan"],
            "poll_date": pd.to_datetime(keep["polldate"], format="%m/%d/%Y", errors="coerce")
                .fillna(pd.to_datetime(keep["polldate"], errors="coerce")).dt.strftime("%Y-%m-%d"),
            "election_date": pd.to_datetime(keep["electiondate"], format="%m/%d/%Y", errors="coerce")
                .fillna(pd.to_datetime(keep["electiondate"], errors="coerce")).dt.strftime("%Y-%m-%d"),
            "days_to_election": keep["time_to_election"].astype(int),
            "sample_size": keep["samplesize"].round().astype(int),
            "dem_margin_poll": keep["margin_poll"].round(2),
            "dem_margin_actual": keep["margin_actual"].round(4),
        }
    ).reset_index(drop=True)


def _two_party_margin(rows: pd.DataFrame) -> tuple[float | None, str]:
    """Top Democrat minus top Republican share (points) and the winner's party."""
    dem = rows.loc[rows["ballot_party"] == "DEM", "percent"].max()
    rep = rows.loc[rows["ballot_party"] == "REP", "percent"].max()
    winner = rows.loc[rows["winner"].astype(str).str.lower() == "true", "ballot_party"]
    winner_party = {"DEM": "Democrat", "REP": "Republican"}.get(winner.iloc[0] if len(winner) else "", "Other")
    margin = None if pd.isna(dem) or pd.isna(rep) else round(float(dem - rep), 2)
    return margin, winner_party


def transform_seats() -> pd.DataFrame:
    """Senate seats on the 3 November 2026 ballot with their last result and 2024 presidential margin."""
    legislators = yaml.safe_load((RAW / "legislators-current.yaml").read_text(encoding="utf-8"))
    senate = pd.read_csv(RAW / "election_results_senate.csv", low_memory=False)
    president = pd.read_csv(RAW / "election_results_presidential.csv", low_memory=False)
    pres = president.loc[(president["cycle"] == 2024) & (president["stage"] == "general")]

    records = []
    for person in legislators:
        term = person["terms"][-1]
        if term["type"] != "sen":
            continue
        # Class II terms end on 3 January 2027; appointees to other seats serve until the
        # 3 November 2026 special election (their roster term ends that day).
        regular = term.get("class") == 2 and term["end"] == "2027-01-03"
        special = term["end"] == ELECTION_DAY_2026
        if not (regular or special):
            continue
        state, seat = term["state"], SEAT_CLASS[term["class"]]
        history = senate.loc[(senate["state_abbrev"] == state) & (senate["office_seat_name"] == seat)]
        history = history.loc[history["stage"].isin(["general", "runoff", "jungle primary"])]
        last_cycle = int(history["cycle"].max())
        latest = history.loc[history["cycle"] == last_cycle]
        stage = "runoff" if (latest["stage"] == "runoff").any() else latest["stage"].iloc[0]
        latest = latest.loc[latest["stage"] == stage]
        if latest["ranked_choice_round"].notna().any():
            latest = latest.loc[latest["ranked_choice_round"] == latest["ranked_choice_round"].min()]
        margin, winner = _two_party_margin(latest)
        pres_margin, _ = _two_party_margin(pres.loc[pres["state_abbrev"] == state])
        party = {"Democrat": "Democrat", "Republican": "Republican"}.get(term["party"], "Independent")
        records.append(
            {
                "state": state,
                "state_name": STATE_NAMES[state],
                "seat_class": seat,
                "election_type": "Special" if special else "Regular",
                "incumbent": person["name"].get("official_full")
                    or f"{person['name']['first']} {person['name']['last']}",
                "incumbent_party": party,
                "incumbent_since": term["start"],
                "last_election_cycle": last_cycle,
                "last_election_stage": stage,
                                "last_winner_party": winner,
                "last_dem_margin": margin,
                "pres_2024_dem_margin": pres_margin,
            }
        )
    seats = pd.DataFrame(records)
    seats["last_election_date"] = seats.apply(_election_date, axis=1)
    return seats.sort_values("state").reset_index(drop=True)


def transform_swing(seats: pd.DataFrame) -> pd.DataFrame:
    """Statewide presidential margins in 2020 and 2024 for the 50 states and DC."""
    president = pd.read_csv(RAW / "election_results_presidential.csv", low_memory=False)
    general = president.loc[(president["stage"] == "general") & president["cycle"].isin([2020, 2024])]
    general = general.loc[general["state_abbrev"].str.fullmatch(r"[A-Z]{2}", na=False) & (general["state_abbrev"] != "PR")]
    records = []
    for state, rows in general.groupby("state_abbrev"):
        margins = {}
        for cycle in (2020, 2024):
            margin, _ = _two_party_margin(rows.loc[rows["cycle"] == cycle])
            margins[cycle] = margin
        records.append({
            "state": state, "state_name": rows["state"].iloc[0],
            "dem_margin_2020": margins[2020], "dem_margin_2024": margins[2024],
            "senate_race_2026": "yes" if state in set(seats["state"]) else "no",
        })
    return pd.DataFrame(records).sort_values("state").reset_index(drop=True)


def _sheet_rows(sheet_name: str, first_header: str) -> list[dict]:
    """Rows under the header row whose first cell is `first_header`, until the first blank first cell.

    Only typed input columns are read; formula columns are ignored and recomputed here, so the
    pipeline never depends on whether Excel recalculated the file before it was saved.
    """
    from openpyxl import load_workbook  # imported here so offline sample mode does not need it

    if not CURATED_XLSX.exists():
        raise FileNotFoundError(f"Curated workbook missing: {CURATED_XLSX}")
    sheet = load_workbook(CURATED_XLSX, read_only=True, data_only=True)[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    start = next((i for i, row in enumerate(rows) if row and row[0] == first_header), None)
    if start is None:
        raise ValueError(f"No header row starting with {first_header!r} on sheet {sheet_name!r}")
    names = [str(v).strip() if v is not None else "" for v in rows[start]]
    records = []
    for number, row in enumerate(rows[start + 1:], start=start + 2):
        if row is None or row[0] is None or str(row[0]).strip() == "":
            break
        record = dict(zip(names, row))
        record["_row"] = number
        records.append(record)
    return records


def _iso(value, where: str) -> str:
    """Accept a typed ISO string or an Excel date cell; reject anything else."""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    try:
        return pd.Timestamp(text).strftime("%Y-%m-%d")
    except (ValueError, TypeError) as error:
        raise ValueError(f"{where}: date {value!r} is not YYYY-MM-DD") from error


def _party(value, where: str) -> str:
    text = str(value or "").strip().upper()
    if text not in {"", "D", "R", "I"}:
        raise ValueError(f"{where}: party {value!r} must be D, R, I or blank")
    return text


def transform_polls_2026() -> pd.DataFrame:
    """2026 Senate polls from the curated workbook's 'Senate polls' sheet."""
    out = []
    for r in _sheet_rows("Senate polls", "State"):
        where = f"Senate polls row {r['_row']}"
        out.append({
            "state": str(r["State code"] or "").strip().upper(),
            "race": str(r["State"]).strip(),
            "seat_status": str(r["Seat"] or "").strip(),
            "candidate_1": str(r["Candidate 1"]).strip(),
            "party_1": _party(r["Party 1"], where),
            "pct_1": r["Candidate 1 (%)"],
            "candidate_2": str(r["Candidate 2"]).strip(),
            "party_2": _party(r["Party 2"], where),
            "pct_2": r["Candidate 2 (%)"],
            "pollster": str(r["Pollster"]).strip(),
            "poll_date": _iso(r["Poll date"], where),
            "sample_note": str(r["Sample note"] or "").strip(),
            "source_url": str(r["Source URL"] or "").strip(),
        })
    frame = pd.DataFrame(out)
    for column in ("pct_1", "pct_2"):
        frame[column] = pd.to_numeric(frame[column], errors="raise").astype(float)
    return frame.replace({"": None})


def transform_generic_2026() -> pd.DataFrame:
    """Generic congressional ballot readings from the curated workbook's 'National' sheet."""
    out = []
    for r in _sheet_rows("National", "Measure"):
        if r.get("Democrats (%)") is None or r.get("Republicans (%)") is None:
            continue   # notes rows such as prediction-market prices
        out.append({
            "measure": str(r["Measure"]).strip(),
            "date": _iso(r["Date"], f"National row {r['_row']}"),
            "dem_pct": float(r["Democrats (%)"]),
            "rep_pct": float(r["Republicans (%)"]),
            "source": str(r["Source"] or "").strip(),
            "source_url": str(r.get("Source URL") or "").strip(),
        })
    return pd.DataFrame(out).replace({"": None})


def _election_date(row: pd.Series) -> str:
    """Federal general election day of the seat's last cycle; Georgia 2020 runoff on 5 January 2021."""
    if row["last_election_stage"] == "runoff" and row["state"] == "GA" and row["last_election_cycle"] == 2020:
        return "2021-01-05"
    year = int(row["last_election_cycle"])
    first = pd.Timestamp(year=year, month=11, day=1)
    monday = first + pd.Timedelta(days=(0 - first.weekday()) % 7)
    return (monday + pd.Timedelta(days=1)).strftime("%Y-%m-%d")


# --------------------------------------------------------------------------- derive
def derive_polls(frame: pd.DataFrame) -> pd.DataFrame:
    """Signed error (positive = poll overstated the Democrat), absolute error and wrong-winner flag."""
    base = ["question_id", "poll_id", "cycle", "race_type", "race", "location", "pollster", "methodology",
            "partisan_sponsor", "poll_date", "election_date", "days_to_election", "sample_size",
            "dem_margin_poll", "dem_margin_actual"]
    frame = frame[base].copy()
    frame["signed_error"] = (frame["dem_margin_poll"] - frame["dem_margin_actual"]).round(4)
    frame["absolute_error"] = frame["signed_error"].abs().round(4)
    frame["called_wrong_winner"] = (
        (frame["dem_margin_poll"] * frame["dem_margin_actual"] < 0)
    ).map({True: "yes", False: "no"})
    return frame.sort_values(["cycle", "race", "poll_date", "question_id"]).reset_index(drop=True)


def derive_seats(frame: pd.DataFrame) -> pd.DataFrame:
    """Absolute size of the last margin and the seat's lean relative to the 2024 presidential vote."""
    base = ["state", "state_name", "seat_class", "election_type", "incumbent", "incumbent_party",
            "incumbent_since", "last_election_cycle", "last_election_stage", "last_election_date",
            "last_winner_party", "last_dem_margin", "pres_2024_dem_margin"]
    frame = frame[base].copy()
    frame["last_margin_size"] = frame["last_dem_margin"].abs().round(2)
    frame["senate_minus_pres_margin"] = (frame["last_dem_margin"] - frame["pres_2024_dem_margin"]).round(2)
    return frame.sort_values("state").reset_index(drop=True)


def derive_swing(frame: pd.DataFrame) -> pd.DataFrame:
    """Swing = 2024 margin minus 2020 margin (negative = moved toward the Republican)."""
    frame = frame[["state", "state_name", "dem_margin_2020", "dem_margin_2024", "senate_race_2026"]].copy()
    frame["swing_2020_2024"] = (frame["dem_margin_2024"] - frame["dem_margin_2020"]).round(2)
    frame["election_date"] = "2024-11-05"
    return frame.sort_values("state").reset_index(drop=True)


def derive_polls_2026(frame: pd.DataFrame) -> pd.DataFrame:
    """Lead of candidate 1, and the Democrat-minus-Republican lead where both parties are stated."""
    base = ["state", "race", "seat_status", "candidate_1", "party_1", "pct_1", "candidate_2", "party_2", "pct_2",
            "pollster", "poll_date", "sample_note", "source_url"]
    frame = frame[base].copy()
    frame["lead_1"] = (frame["pct_1"] - frame["pct_2"]).round(1)
    dem_first = (frame["party_1"] == "D") & (frame["party_2"] == "R")
    rep_first = (frame["party_1"] == "R") & (frame["party_2"] == "D")
    frame["dem_lead"] = None
    frame.loc[dem_first, "dem_lead"] = frame.loc[dem_first, "lead_1"]
    frame.loc[rep_first, "dem_lead"] = -frame.loc[rep_first, "lead_1"]
    frame["dem_lead"] = pd.to_numeric(frame["dem_lead"]).round(1)
    return frame.sort_values(["state", "poll_date", "pollster"]).reset_index(drop=True)


def derive_generic_2026(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame[["measure", "date", "dem_pct", "rep_pct", "source", "source_url"]].copy()
    frame["dem_lead"] = (frame["dem_pct"] - frame["rep_pct"]).round(1)
    return frame.sort_values(["date", "measure"]).reset_index(drop=True)


DERIVE = {"polls": derive_polls, "seats": derive_seats, "swing": derive_swing,
          "polls_2026": derive_polls_2026, "generic_2026": derive_generic_2026}
CURATED = {"polls_2026": transform_polls_2026, "generic_2026": transform_generic_2026}


# --------------------------------------------------------------------------- validate
def load_contract() -> dict:
    manifest = yaml.safe_load(MANIFEST.read_text(encoding="utf-8"))
    return {dataset["id"]: dataset for dataset in manifest["datasets"]}


def read_dataset(path: Path, spec: dict) -> pd.DataFrame:
    """Read a CSV with declared string columns kept as text; only empty fields are missing."""
    strings = {name: str for name, column in spec["columns"].items() if column["type"] in {"string", "date"}}
    return pd.read_csv(path, dtype=strings, keep_default_na=False, na_values=[""])


def validate(frame: pd.DataFrame, spec: dict) -> None:
    """Raise ValueError if the frame breaks any expectation declared in the manifest."""
    name, columns = spec["id"], spec["columns"]
    problems = []
    missing = set(columns).difference(frame.columns)
    if missing:
        raise ValueError(f"{name}: missing columns {sorted(missing)}")
    extra = set(frame.columns).difference(columns)
    if extra:
        problems.append(f"undeclared columns {sorted(extra)}")
    if len(frame) < spec.get("minimum_rows", 1):
        problems.append(f"only {len(frame)} rows; at least {spec['minimum_rows']} declared")
    keys = spec["primary_key"]
    if frame.duplicated(keys).any():
        problems.append(f"duplicate primary keys {keys}: {int(frame.duplicated(keys).sum())}")
    for column, rule in columns.items():
        series = frame[column]
        present = series.dropna()
        if not rule["nullable"] and series.isna().any():
            problems.append(f"{column}: {int(series.isna().sum())} missing values in a non-nullable column")
        if rule["type"] in {"integer", "number"}:
            numeric = pd.to_numeric(present, errors="coerce")
            if numeric.isna().any():
                problems.append(f"{column}: non-numeric values")
                continue
            if rule["type"] == "integer" and not (numeric % 1 == 0).all():
                problems.append(f"{column}: non-integer values")
            if "minimum" in rule and (numeric < rule["minimum"]).any():
                problems.append(f"{column}: values below {rule['minimum']} {rule.get('unit', '')}")
            if "maximum" in rule and (numeric > rule["maximum"]).any():
                problems.append(f"{column}: values above {rule['maximum']} {rule.get('unit', '')}")
        if rule["type"] == "date":
            dates = pd.to_datetime(present, errors="coerce", format="%Y-%m-%d")
            if dates.isna().any():
                problems.append(f"{column}: dates not in ISO YYYY-MM-DD form")
            if "date_minimum" in rule and (dates < pd.Timestamp(rule["date_minimum"])).any():
                problems.append(f"{column}: dates before {rule['date_minimum']}")
            if "date_maximum" in rule and (dates > pd.Timestamp(rule["date_maximum"])).any():
                problems.append(f"{column}: dates after {rule['date_maximum']}")
        if "allowed_values" in rule and not present.isin(rule["allowed_values"]).all():
            unexpected = sorted(set(present) - set(rule["allowed_values"]))
            problems.append(f"{column}: unexpected values {unexpected[:5]}")
    if spec.get("date_column") and spec.get("latest_observation"):
        latest = pd.to_datetime(frame[spec["date_column"]]).max().strftime("%Y-%m-%d")
        if latest != spec["latest_observation"]:
            problems.append(f"latest observation {latest} differs from manifest {spec['latest_observation']}")
    if problems:
        raise ValueError(f"{name} failed validation: " + "; ".join(problems))


def check_curated(frames: dict[str, pd.DataFrame]) -> None:
    """Rules for hand-entered rows that a column contract cannot express."""
    problems = []
    today = pd.Timestamp.now(tz="UTC").strftime("%Y-%m-%d")
    for name in ("polls_2026", "generic_2026"):
        frame = frames.get(name)
        if frame is None:
            continue
        bad_url = ~frame["source_url"].fillna("").str.match(r"^https?://")
        if bad_url.any():
            problems.append(f"{name}: {int(bad_url.sum())} row(s) without an http(s) source URL")
        date_col = "poll_date" if name == "polls_2026" else "date"
        if (frame[date_col] > today).any():
            problems.append(f"{name}: dates in the future: {sorted(frame.loc[frame[date_col] > today, date_col])}")
    polls = frames.get("polls_2026")
    seats = frames.get("seats")
    if polls is not None and seats is not None:
        unknown = sorted(set(polls["state"]) - set(seats["state"]))
        if unknown:
            problems.append(f"polls_2026: state codes with no 2026 Senate contest: {unknown}")
        total = polls["pct_1"] + polls["pct_2"]
        if (total > 100).any():
            problems.append("polls_2026: candidate shares add up to more than 100% in some rows")
    if problems:
        raise ValueError("Curated data failed validation: " + "; ".join(problems))


def _cross_checks(frames: dict[str, pd.DataFrame]) -> None:
    """Checks spanning rows: every seat on the ballot appears once, and errors reproduce."""
    seats = frames["seats"]
    if len(seats) != 35:
        raise ValueError(f"Expected 35 Senate contests on the 2026 ballot (33 Class II + 2 specials), found {len(seats)}")
    swing = frames["swing"]
    if len(swing) != 51:
        raise ValueError(f"Expected 50 states plus DC in the swing table, found {len(swing)}")
    if set(swing.loc[swing["senate_race_2026"] == "yes", "state"]) != set(seats["state"]):
        raise ValueError("Swing table's 2026 Senate flag does not match the seats dataset")
    check_curated(frames)
    polls = frames["polls"]
    recomputed = (polls["dem_margin_poll"] - polls["dem_margin_actual"]).abs()
    if (recomputed - polls["absolute_error"]).abs().max() > 1e-3:
        raise ValueError("absolute_error does not reproduce |poll margin - actual margin|")


# --------------------------------------------------------------------------- main
def build(sample_mode: bool) -> dict[str, pd.DataFrame]:
    contract = load_contract()
    if sample_mode:
        base = {name: read_dataset(ROOT / files[0], contract[name]) for name, files in DATASET_FILES.items()}
    else:
        if not RAW.exists() or not all((RAW / name).exists() for name in RAW_SOURCES):
            acquire()
        seats = transform_seats()
        base = {"polls": transform_polls(), "seats": seats, "swing": transform_swing(seats),
                **{name: transform() for name, transform in CURATED.items()}}
    frames = {name: DERIVE[name](frame) for name, frame in base.items()}
    for name, frame in frames.items():
        validate(frame, contract[name])
    _cross_checks(frames)
    return frames


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--acquire", action="store_true", help="Refresh data/raw/ and stop")
    parser.add_argument("--validate", action="store_true", help="Validate the processed outputs and stop")
    parser.add_argument("--refresh-sample", action="store_true", help="Live mode: also rewrite the bundled samples")
    args = parser.parse_args()
    sample_mode = os.getenv("DASHBOARD_SAMPLE_MODE", "0") == "1"
    if args.acquire or args.refresh_sample:
        if sample_mode:
            parser.error("Unset DASHBOARD_SAMPLE_MODE before acquiring live data")
    if args.acquire:
        acquire()
        return
    if args.validate:
        contract = load_contract()
        for name, (_, processed) in DATASET_FILES.items():
            validate(read_dataset(ROOT / processed, contract[name]), contract[name])
        print("Processed data valid")
        return
    frames = build(sample_mode)
    for name, frame in frames.items():
        sample, processed = DATASET_FILES[name]
        targets = [processed] + ([sample] if args.refresh_sample else [])
        for target in targets:
            path = ROOT / target
            path.parent.mkdir(parents=True, exist_ok=True)
            frame.to_csv(path, index=False)
            print(f"Validated and wrote {len(frame):,} rows to {target}")


if __name__ == "__main__":
    main()
